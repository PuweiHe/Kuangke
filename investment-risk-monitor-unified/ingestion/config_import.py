"""CSV/XLSX configuration import with bounded input and atomic upserts."""
import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path

from core.holding_ratio import account_scope

SCHEMAS = {
    'rules': ('risk_rule_config', ('monitor_id',),
              ('monitor_id', 'monitor_type', 'monitor_item', 'monitor_title', 'threshold', 'scope', 'is_enabled')),
    'dimensions': ('risk_dimension_config', ('monitor_id', 'dimension_code'),
                   ('monitor_id', 'dimension_code', 'trust_dimension')),
    'assets': ('risk_assets_list', ('category', 'category_2nd', 'asset_code', 'year'),
               ('category', 'category_2nd', 'asset_code', 'year', 'is_valid')),
}


class ImportValidationError(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__('; '.join(errors))


def read_rows(path, sheet=None):
    """Reject formulas rather than trusting potentially stale cached results."""
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            reader = csv.reader(stream)
            headers = next(reader, [])
            rows = list(reader)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False, keep_links=True)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ImportValidationError(['Select a sheet explicitly for a multi-sheet workbook'])
            ws = book[sheet] if sheet else book.active
            rows = []
            for index, cells in enumerate(ws.iter_rows()):
                if index > 10000:
                    raise ImportValidationError(['Input exceeds 10000 data rows'])
                if any(c.data_type in {'f', 'e'} for c in cells):
                    raise ImportValidationError([f'Row {index + 1}: formulas and Excel errors are not configuration values'])
                rows.append([c.value for c in cells])
            headers = rows.pop(0) if rows else []
        finally:
            book.close()
    else:
        raise ImportValidationError(['Only .csv and .xlsx are supported'])
    if len(rows) > 10000:
        raise ImportValidationError(['Input exceeds 10000 data rows'])
    if not headers or any(not isinstance(h, str) or not h.strip() for h in headers):
        raise ImportValidationError(['Headers must be nonempty text'])
    headers = [h.strip() for h in headers]
    if len(set(headers)) != len(headers):
        raise ImportValidationError(['Duplicate column names'])
    records = []
    for index, row in enumerate(rows, 2):
        if all(v is None or v == '' for v in row):
            continue
        if len(row) != len(headers):
            raise ImportValidationError([f'Row {index}: wrong column count'])
        records.append((index, dict(zip(headers, row))))
    return headers, records


def validate(kind, headers, records):
    if kind not in SCHEMAS:
        raise ImportValidationError(['Unknown configuration kind'])
    _, keys, columns = SCHEMAS[kind]
    if set(headers) != set(columns):
        raise ImportValidationError(['Columns must match: ' + ', '.join(columns)])
    errors, output, seen = [], [], set()
    for row_number, original in records:
        try:
            row = {}
            for field in columns:
                value = original[field]
                if field in keys and field != 'year' and not isinstance(value, str):
                    raise ValueError(f'{field} must be text; preserve identifier leading zeroes')
                row[field] = str(value).strip() if value is not None else ''
                if row[field].startswith(('=', '+', '@')) or row[field] in {'#REF!', '#N/A', '#VALUE!'}:
                    raise ValueError(f'{field} contains an expression or spreadsheet error')
                limit = 500 if field == 'scope' else 255
                if field in {'monitor_id', 'dimension_code', 'trust_dimension'}: limit = 50
                if field in {'category', 'category_2nd', 'monitor_type'}: limit = 50
                if field == 'asset_code': limit = 20
                if len(row[field]) > limit:
                    raise ValueError(f'{field} is too long')
                if field not in {'scope', 'threshold'} and not row[field]:
                    raise ValueError(f'{field} is required')
            for field in ('is_enabled', 'is_valid'):
                if field in row:
                    if row[field] not in {'0', '1'}:
                        raise ValueError(f'{field} must be 0 or 1')
                    row[field] = int(row[field])
            if kind == 'rules':
                if row['scope']:
                    account_scope(row['scope'])
                raw = row['threshold']
                if not raw:
                    raise ValueError('threshold is required in the numeric-rule template')
                number = Decimal(raw[:-1]) / 100 if raw.endswith('%') else Decimal(raw)
                if not number.is_finite() or number < 0:
                    raise ValueError('threshold must be finite and nonnegative')
                row['threshold'] = str(number)
                row['scope'] = row['scope'] or None
            if kind == 'assets':
                if not row['year'].isdigit() or not 1900 <= int(row['year']) <= 9999:
                    raise ValueError('year must be a four-digit year')
                row['year'] = int(row['year'])
            key = tuple(row[k] for k in keys)
            if key in seen:
                raise ValueError('Duplicate natural key')
            seen.add(key)
            output.append(row)
        except (ValueError, InvalidOperation, TypeError) as exc:
            # Report field/row diagnostics, never the original row payload.
            errors.append(f'Row {row_number}: {exc if not isinstance(exc, InvalidOperation) else "Invalid numeric threshold"}')
    if errors:
        raise ImportValidationError(errors)
    if not output:
        raise ImportValidationError(['No data rows'])
    return output


def prepare(kind, path, sheet=None):
    return validate(kind, *read_rows(path, sheet))


def upsert_sql(kind, dialect):
    if dialect not in {'mysql', 'pg'}:
        raise ValueError('Unsupported database dialect')
    table, keys, columns = SCHEMAS[kind]
    updates = [c for c in columns if c not in keys]
    suffix = ('ON DUPLICATE KEY UPDATE ' + ', '.join(f'{c}=VALUES({c})' for c in updates)
              if dialect == 'mysql' else
              'ON CONFLICT (' + ', '.join(keys) + ') DO UPDATE SET ' + ', '.join(f'{c}=EXCLUDED.{c}' for c in updates))
    return f'INSERT INTO {table} ({", ".join(columns)}) VALUES ({", ".join(["%s"] * len(columns))}) {suffix}'


def apply_import(connection, dialect, kind, rows):
    """Own one transaction. Call with an exclusive, non-autocommit connection."""
    columns = SCHEMAS[kind][2]
    rows = validate(kind, list(columns), [(i, r) for i, r in enumerate(rows, 2)])
    try:
        with connection.cursor() as cursor:
            if kind == 'dimensions':
                for monitor_id in sorted({r['monitor_id'] for r in rows}):
                    cursor.execute('SELECT monitor_id FROM risk_rule_config WHERE monitor_id = %s', (monitor_id,))
                    if cursor.fetchone() is None:
                        raise ImportValidationError(['Dimension refers to an unknown monitor'])
            cursor.executemany(upsert_sql(kind, dialect), [tuple(row[c] for c in columns) for row in rows])
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return len(rows)
